from docx import Document

# Ouvrir le doc
doc = Document(r'project\docsTechnique\DOCUMENT_TECHNIQUE.docx')

# Afficher le contenu pour voir la structure
print("=== DOCUMENT CONTENT (first 40 paragraphs) ===")
for i, para in enumerate(doc.paragraphs[:40]):
    style = para.style.name if para.style else "None"
    print(f"{i}: [{style}] {para.text[:90]}")

print("\n\n=== TABLES IN DOC ===")
for i, table in enumerate(doc.tables):
    print(f"Table {i}: {len(table.rows)} rows, {len(table.columns)} cols")
    print(f"  Content: {table.rows[0].cells[0].text[:50] if table.rows else 'empty'}")
